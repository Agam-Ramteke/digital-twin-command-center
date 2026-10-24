from __future__ import annotations

import re
import textwrap
import zipfile
from collections import defaultdict
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\ASUS\OneDrive\Desktop\digital-twin-command-center")
WORK = ROOT / "thesis_work"
TEMPLATE = WORK / "template-copy.docx"
CONTENT = WORK / "thesis_content.md"
OUTPUT = ROOT / "Digital_Twin_Orchestration_Platform_Thesis.docx"
FIGURES = WORK / "figures"
SCREENSHOT = ROOT / "browser_screenshot_selected.png"
EXPECTED_TEMPLATE_SHA = "99083E5CC6085CB742016C557642D00BD5FEB490F0E9BC235E27E0A9813023E6"

INK = (24, 34, 45)
MID = (75, 88, 101)
LIGHT = (235, 240, 244)
WHITE = (255, 255, 255)
ACCENT = (45, 105, 145)


def verify_template() -> None:
    import hashlib
    digest = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest().upper()
    if digest != EXPECTED_TEMPLATE_SHA:
        raise RuntimeError(f"Template working copy changed: {digest}")


def get_font(size: int, bold: bool = False, mono: bool = False):
    if mono:
        candidates = [r"C:\Windows\Fonts\consola.ttf", r"C:\Windows\Fonts\lucon.ttf"]
    elif bold:
        candidates = [r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\calibrib.ttf"]
    else:
        candidates = [r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\calibri.ttf"]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def wrap_text(text: str, font, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        candidate = word if not line else f"{line} {word}"
        if ImageDraw.Draw(Image.new("RGB", (2, 2))).textlength(candidate, font=font) <= max_width:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def make_canvas(title: str, subtitle: str = ""):
    image = Image.new("RGB", (1600, 900), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), title, font=get_font(34, True), fill=INK)
    if subtitle:
        draw.text((72, 82), subtitle, font=get_font(21), fill=MID)
    draw.line((70, 125, 1530, 125), fill=(170, 180, 188), width=2)
    return image, draw


def draw_box(draw, xy, title: str, body: str = "", fill=LIGHT, outline=INK, title_font=None, body_font=None):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=15, fill=fill, outline=outline, width=3)
    title_font = title_font or get_font(23, True)
    body_font = body_font or get_font(18)
    title_lines = wrap_text(title, title_font, x2 - x1 - 28)
    body_lines = wrap_text(body, body_font, x2 - x1 - 28)
    total_h = len(title_lines) * 30 + len(body_lines) * 25
    y = y1 + max(14, ((y2 - y1) - total_h) // 2)
    for line in title_lines:
        width = draw.textlength(line, font=title_font)
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=title_font, fill=INK)
        y += 30
    if body_lines:
        y += 5
    for line in body_lines:
        width = draw.textlength(line, font=body_font)
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=body_font, fill=MID)
        y += 25


def arrow(draw, start, end, color=ACCENT, width=5):
    import math
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    size = 15
    p1 = (x2, y2)
    p2 = (x2 - size * math.cos(angle - 0.45), y2 - size * math.sin(angle - 0.45))
    p3 = (x2 - size * math.cos(angle + 0.45), y2 - size * math.sin(angle + 0.45))
    draw.polygon((p1, p2, p3), fill=color)


def save_figure(image, name: str):
    FIGURES.mkdir(parents=True, exist_ok=True)
    path = FIGURES / f"{name}.png"
    image.save(path, optimize=True)
    return path


def create_diagrams() -> None:
    # Service and data architecture.
    im, d = make_canvas("Digital Twin Orchestration Platform", "Separated model, state, process, scenario, and presentation services")
    draw_box(d, (80, 195, 360, 365), "Five-station line", "Causal simulator or future edge adapter\nTelemetry and event envelopes")
    draw_box(d, (455, 195, 730, 365), "MQTT broker", "Canonical factory/v1 topics\nTCP and WebSocket listeners")
    draw_box(d, (835, 170, 1130, 390), "Live Twin service", "Validate, order, reduce, measure timing\nCurrent state plus event history")
    draw_box(d, (1250, 170, 1515, 390), "MongoDB / memory", "Twins, events, telemetry, tokens,\nscenario records")
    arrow(d, (360, 280), (455, 280))
    arrow(d, (730, 280), (835, 280))
    arrow(d, (1130, 280), (1250, 280))
    draw_box(d, (80, 545, 300, 720), "React dashboard", "Live cockpit, production,\nscenarios, events, system")
    draw_box(d, (430, 545, 700, 720), "FastAPI", "Versioned REST routes,\nhealth and readiness")
    draw_box(d, (830, 545, 1100, 720), "Process Twin", "DAG, buffers, token lineage,\nbottleneck and defect trace")
    draw_box(d, (1230, 545, 1500, 720), "Scenario Worker", "Frozen snapshot; same-seed\nbaseline and counterfactual")
    arrow(d, (300, 632), (430, 632))
    arrow(d, (700, 632), (830, 632))
    arrow(d, (1100, 632), (1230, 632))
    arrow(d, (965, 390), (965, 545))
    arrow(d, (1080, 390), (1365, 545))
    save_figure(im, "architecture")

    im, d = make_canvas("Event and state data flow", "The same reducer handles broker telemetry and controlled HTTP ingestion")
    xs = [65, 355, 655, 955, 1250]
    labels = [
        ("Producer", "Station observation\nsource and UTC time"),
        ("Envelope", "Stable ID, sequence,\nschema and payload"),
        ("Transport", "MQTT subscriber or\nHTTP ingestion route"),
        ("Reducer", "Validate and compare\nsequence and identity"),
        ("Materialized state", "Current twin plus\nimmutable event history"),
    ]
    for x, (title, body) in zip(xs, labels):
        draw_box(d, (x, 285, x + 230, 500), title, body)
    for i in range(4):
        arrow(d, (xs[i] + 230, 390), (xs[i + 1], 390))
    draw_box(d, (420, 640, 720, 790), "Duplicate event", "Identity barrier; state unchanged", fill=(247, 247, 247))
    draw_box(d, (850, 640, 1150, 790), "Late sequence", "Retain event; current state unchanged", fill=(247, 247, 247))
    arrow(d, (1060, 500), (570, 640), color=MID, width=3)
    arrow(d, (1060, 500), (1000, 640), color=MID, width=3)
    save_figure(im, "dataflow")

    im, d = make_canvas("Live Twin reducer decision", "Current state advances only for a valid, non-duplicate, newer observation")
    draw_box(d, (585, 175, 1005, 280), "Receive telemetry event", "Capture receipt time at consumer boundary", fill=(230, 238, 245))
    draw_box(d, (585, 345, 1005, 450), "Validate envelope and payload", "Kind, UTC timestamps, station fields, machine identity")
    arrow(d, (795, 280), (795, 345))
    draw_box(d, (585, 515, 1005, 620), "Check event identity", "Unique event key at repository boundary")
    arrow(d, (795, 450), (795, 515))
    draw_box(d, (100, 700, 445, 825), "Duplicate", "No new state or event write", fill=(244, 244, 244))
    draw_box(d, (625, 700, 970, 825), "Late or equal sequence", "Retain event; do not roll state back", fill=(244, 244, 244))
    draw_box(d, (1150, 700, 1495, 825), "Newer sequence", "Reduce state, increment revision,\nstore measured timing", fill=(230, 238, 245))
    arrow(d, (680, 620), (270, 700), color=MID, width=3)
    arrow(d, (795, 620), (795, 700), color=MID, width=3)
    arrow(d, (910, 620), (1320, 700), color=ACCENT, width=4)
    save_figure(im, "reducer")

    im, d = make_canvas("Five-station material flow", "Sequential pipeline with a maximum inter-station buffer of five units")
    stations = [
        ("STAMPING-01", "8.5 s"),
        ("CNC-01", "42.0 s"),
        ("WELDING-01", "15.0 s"),
        ("INSPECTION-01", "6.0 s"),
        ("PACKAGING-01", "10.0 s"),
    ]
    start = 40
    width = 235
    gap = 75
    for i, (name, cycle) in enumerate(stations):
        x = start + i * (width + gap)
        draw_box(d, (x, 300, x + width, 455), name, f"Configured cycle: {cycle}", fill=(230, 238, 245) if i == 1 else LIGHT)
        if i < 4:
            ax = x + width
            arrow(d, (ax, 378), (ax + gap, 378))
            d.rounded_rectangle((ax + 3, 500, ax + gap - 3, 570), radius=10, fill=(245, 245, 245), outline=MID, width=2)
            d.text((ax + 20, 505), "WIP", font=get_font(13, True), fill=INK)
            d.text((ax + 25, 530), "≤ 5", font=get_font(15, True), fill=INK)
    d.text((110, 665), "Downstream starvation guard: no input part available -> station waits.", font=get_font(22, True), fill=INK)
    d.text((110, 715), "Upstream blocking guard: downstream buffer full -> upstream station pauses.", font=get_font(22, True), fill=INK)
    d.text((110, 785), "Invariant: Packaging <= Inspection <= Welding <= CNC <= Stamping output.", font=get_font(23, True), fill=ACCENT)
    save_figure(im, "materialflow")

    im, d = make_canvas("Snapshot-based counterfactual execution", "Baseline and intervention use the same immutable starting snapshot and seed")
    draw_box(d, (590, 165, 1010, 290), "Frozen line snapshot", "Ordered current state for five stations", fill=(230, 238, 245))
    draw_box(d, (95, 390, 625, 545), "Baseline factory", "Same initial state and seed\nNo intervention override")
    draw_box(d, (975, 390, 1505, 545), "Counterfactual factory", "Same initial state and seed\nApply selected station override")
    arrow(d, (700, 290), (360, 390))
    arrow(d, (900, 290), (1240, 390))
    draw_box(d, (95, 640, 625, 790), "Baseline outcome", "Output, scrap, quality, energy,\nhealth, bottleneck, trajectory")
    draw_box(d, (975, 640, 1505, 790), "Counterfactual outcome", "Output, scrap, quality, energy,\nhealth, bottleneck, trajectory")
    arrow(d, (360, 545), (360, 640))
    arrow(d, (1240, 545), (1240, 640))
    draw_box(d, (620, 635, 980, 800), "Compare and verify", "Compute deltas; re-check Live Twin\nfingerprint; persist result", fill=(230, 238, 245))
    arrow(d, (625, 715), (620, 715))
    arrow(d, (975, 715), (980, 715))
    save_figure(im, "scenario")

    im, d = make_canvas("Container deployment topology", "Local stack and Kubernetes resources use the same service boundaries")
    draw_box(d, (95, 210, 400, 385), "Browser", "React application\nHTTP and MQTT WebSocket")
    draw_box(d, (640, 210, 945, 385), "FastAPI backend", "Simulator, twin, process,\nscenario, health and readiness")
    draw_box(d, (1180, 150, 1500, 320), "Mosquitto", "TCP 1883\nWebSocket 9001")
    draw_box(d, (1180, 430, 1500, 600), "MongoDB", "Twins, events, tokens,\nscenario persistence")
    arrow(d, (400, 275), (640, 275))
    arrow(d, (945, 250), (1180, 235))
    arrow(d, (945, 335), (1180, 500))
    draw_box(d, (180, 640, 700, 805), "Docker Compose", "Four services, health checks, shared network,\npersistent MongoDB volume")
    draw_box(d, (900, 640, 1420, 805), "Kubernetes", "Namespace, ConfigMap, StatefulSet, Deployments,\nServices, Ingress, probes and Kustomize")
    save_figure(im, "deployment")

    def bar_chart(name, title, subtitle, labels, values, x_max, axis_label, value_formatter):
        image, draw = make_canvas(title, subtitle)
        left, right, top, bottom = 320, 1480, 190, 760
        draw.line((left, top, left, bottom), fill=INK, width=3)
        draw.line((left, bottom, right, bottom), fill=INK, width=3)
        for tick in range(0, int(x_max) + 1, max(1, int(x_max // 5))):
            x = left + (right - left) * tick / x_max
            draw.line((x, top, x, bottom), fill=(220, 225, 230), width=2)
            draw.text((x - 16, bottom + 15), str(tick), font=get_font(18), fill=MID)
        bar_h = 76
        spacing = 104
        for i, (label, value) in enumerate(zip(labels, values)):
            y = top + 25 + i * spacing
            draw.text((left - 240, y + 20), label, font=get_font(22, True), fill=INK)
            w = (right - left) * value / x_max
            draw.rounded_rectangle((left + 1, y, left + w, y + bar_h), radius=8, fill=ACCENT if i == 1 else (105, 130, 150))
            draw.text((left + w + 12, y + 23), value_formatter(value), font=get_font(20, True), fill=INK)
        draw.text((left, 815), axis_label, font=get_font(20), fill=MID)
        save_figure(image, name)

    bar_chart(
        "cyclechart", "Configured station cycle times",
        "Source: DEFAULT_STATIONS in the project model; configuration values, not measured cycle times",
        ["Stamping", "CNC", "Welding", "Inspection", "Packaging"],
        [8.5, 42.0, 15.0, 6.0, 10.0], 50, "Configured seconds per cycle",
        lambda v: f"{v:.1f} s")

    bar_chart(
        "capacitychart", "Idealized configured rate comparison",
        "Derived as 3,600 / nominal cycle time; excludes downtime, yield, queues, and shift effects",
        ["Stamping", "CNC", "Welding", "Inspection", "Packaging"],
        [3600/8.5, 3600/42.0, 3600/15.0, 3600/6.0, 3600/10.0], 700,
        "Theoretical cycles per hour from configured cycle time",
        lambda v: f"{v:.1f}")

    image, draw = make_canvas("Repository-reported final verification summary", "Source: BUILD_PROGRESS.md, material-flow redesign entry dated 2026-10-20")
    left, right, top, bottom = 440, 1280, 235, 625
    for tick in (0, 12, 24, 36, 48):
        x = left + (right - left) * tick / 48
        draw.line((x, top, x, bottom), fill=(220, 225, 230), width=2)
        draw.text((x - 12, bottom + 12), str(tick), font=get_font(18), fill=MID)
    rows = [
        ("Passed", 46, (75, 115, 140)),
        ("Skipped", 2, (202, 151, 70)),
    ]
    for label, value, color in rows:
        y = top + (55 if label == "Passed" else 220)
        draw.text((150, y + 16), label, font=get_font(25, True), fill=INK)
        width = max(12, (right - left) * value / 48)
        draw.rounded_rectangle((left, y, left + width, y + 64), radius=8, fill=color)
        draw.text((left + width + 14, y + 16), str(value), font=get_font(24, True), fill=INK)
    draw.line((left, bottom, right, bottom), fill=INK, width=3)
    draw.text((left, 680), "Recorded backend cases (48 total)", font=get_font(20), fill=MID)
    draw.text((150, 745), "Run marked OK; two skipped; frontend production build passed.", font=get_font(21), fill=INK)
    draw.text((150, 790), "Repository evidence only; not a physical factory or performance benchmark.", font=get_font(20), fill=MID)
    save_figure(image, "testchart")


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, name="Times New Roman", size=12, bold=None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run()
    set_run_font(r, size=9, color="555555")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        r._r.append(node)


def add_field(paragraph, instruction: str, placeholder: str):
    simple = OxmlElement("w:fldSimple")
    simple.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = placeholder
    run.append(text)
    simple.append(run)
    paragraph._p.append(simple)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_table_header_repeat(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_font(cell, size=8.8, bold=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold)


def format_table(table, headers: bool = True):
    table.style = "Table Grid"
    table.autofit = True
    table.allow_autofit = True
    size = 7.1 if len(table.columns) >= 12 else (8.0 if len(table.columns) >= 8 else 9.0)
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        if headers and row_index == 0:
            add_table_header_repeat(row)
        for cell in row.cells:
            if headers and row_index == 0:
                shade_cell(cell, "DDE5EB")
            set_cell_font(cell, size=size, bold=(headers and row_index == 0))


def add_md_table(doc, rows: list[list[str]]):
    if not rows:
        return
    columns = max(len(r) for r in rows)
    table = doc.add_table(rows=1, cols=columns)
    for i, value in enumerate(rows[0]):
        table.rows[0].cells[i].text = value
    for values in rows[1:]:
        cells = table.add_row().cells
        for i in range(columns):
            cells[i].text = values[i] if i < len(values) else ""
    format_table(table, headers=True)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)


def add_body_paragraph(doc, text: str, front: bool = False):
    text = text.replace("**", "").strip()
    if not text:
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.28)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15 if front else 1.35
    paragraph.paragraph_format.widow_control = True
    run = paragraph.add_run(text)
    set_run_font(run, size=11 if front else 12)


FIGURE_TITLES = {
    "architecture": "Reference architecture of the Digital Twin Orchestration Platform",
    "dataflow": "Telemetry event and state reduction data flow",
    "reducer": "Live Twin reducer decision flow",
    "materialflow": "Five-station material pipeline and flow constraints",
    "scenario": "Frozen snapshot baseline and counterfactual execution",
    "deployment": "Container and Kubernetes deployment topology",
    "cyclechart": "Configured station cycle times",
    "capacitychart": "Idealized capacity derived from configured cycle times",
    "testchart": "Repository-reported final verification summary",
}


def add_figure(doc, name: str, chapter_number: int, figure_counts):
    path = SCREENSHOT if name == "screenshot" else (FIGURES / f"{name}.png")
    if not path.exists():
        raise FileNotFoundError(path)
    figure_counts[chapter_number] += 1
    figure_no = f"{chapter_number}.{figure_counts[chapter_number]}"
    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.keep_with_next = True
    image_p.paragraph_format.space_before = Pt(5)
    image_p.paragraph_format.space_after = Pt(0)
    image_run = image_p.add_run()
    width = Inches(6.65 if name == "screenshot" else 6.7)
    image_run.add_picture(str(path), width=width)
    drawing = image_run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", "Illustrative project figure. Simulator chart values are model configuration or logged software verification, as stated in the caption.")
    caption = doc.add_paragraph()
    caption.style = "Caption"
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(9)
    prefix = caption.add_run(f"Figure {chapter_number}.")
    set_run_font(prefix, size=9.5, italic=True)
    sequence = OxmlElement("w:fldSimple")
    if figure_counts[chapter_number] == 1:
        sequence.set(qn("w:instr"), "SEQ Figure \\r 1 \\* ARABIC")
    else:
        sequence.set(qn("w:instr"), "SEQ Figure \\n \\* ARABIC")
    seq_run = OxmlElement("w:r")
    seq_text = OxmlElement("w:t")
    seq_text.text = str(figure_counts[chapter_number])
    seq_run.append(seq_text)
    sequence.append(seq_run)
    caption._p.append(sequence)
    suffix = caption.add_run(f": {FIGURE_TITLES.get(name, 'Prototype dashboard capture; displayed readings are simulator state, not physical factory measurements.')}")
    set_run_font(suffix, size=9.5, italic=True)


def add_contents_pages(doc):
    p = doc.add_paragraph()
    r = p.add_run("TABLE OF CONTENTS")
    set_run_font(r, size=16, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Update this table in Word.")

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LIST OF FIGURES")
    set_run_font(r, size=16, bold=True)
    add_field(doc.add_paragraph(), 'TOC \\h \\z \\c "Figure"', "Update this table in Word.")

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ABBREVIATIONS")
    set_run_font(r, size=16, bold=True)
    rows = [
        ["Term", "Meaning"],
        ["API", "Application Programming Interface"],
        ["CPS", "Cyber-Physical System"],
        ["DAG", "Directed Acyclic Graph"],
        ["IIoT", "Industrial Internet of Things"],
        ["KPI", "Key Performance Indicator"],
        ["MQTT", "Message Queuing Telemetry Transport"],
        ["OEM", "Original Equipment Manufacturer"],
        ["OPC UA", "OPC Unified Architecture"],
        ["REST", "Representational State Transfer"],
        ["WIP", "Work in Process"],
    ]
    add_md_table(doc, rows)


def add_cover(doc, logo_bytes: bytes):
    section = doc.sections[0]
    border = OxmlElement("w:pgBorders")
    border.set(qn("w:offsetFrom"), "page")
    border.set(qn("w:display"), "firstPage")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "12")
        element.set(qn("w:space"), "18")
        element.set(qn("w:color"), "000000")
        border.append(element)
    section._sectPr.append(border)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("DESIGN AND IMPLEMENTATION OF A DIGITAL TWIN\nORCHESTRATION PLATFORM FOR INDUSTRIAL IoT MANUFACTURING")
    set_run_font(r, size=17, bold=True)

    for line, size, bold in [
        ("A Project Report Submitted in Partial Fulfilment of the Requirements for the Degree of", 11.5, False),
        ("Bachelor of Technology", 15, True),
        ("in", 11.5, False),
        ("Industrial IoT", 14, True),
        ("to", 11.5, False),
        ("PRIYADARSHINI COLLEGE OF ENGINEERING", 15, True),
        ("An Autonomous Institute Affiliated to Rashtrasant Tukdoji Maharaj Nagpur University, Nagpur", 10.5, False),
        ("Under the Guidance of", 12, False),
        ("[Project Guide Name and Designation]", 12, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5 if bold else 2)
        r = p.add_run(line)
        set_run_font(r, size=size, bold=bold)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("Submitted by")
    set_run_font(r, size=11, italic=True)

    names = doc.add_table(rows=3, cols=2)
    names.style = "Table Grid"
    names.autofit = False
    for idx, label in enumerate([
        "[Student Name 1]", "[Student Name 2]", "[Student Name 3]",
        "[Student Name 4]", "[Student Name 5]", "[Add or remove student fields as required]",
    ]):
        cell = names.rows[idx // 2].cells[idx % 2]
        cell.text = label
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=10.5, bold=True)
    names.autofit = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(7)
    p.add_run().add_picture(BytesIO(logo_bytes), width=Inches(1.25))
    for line, size, bold in [
        ("DEPARTMENT OF INDUSTRIAL IoT", 13, True),
        ("PRIYADARSHINI COLLEGE OF ENGINEERING, NAGPUR-440019", 11, False),
        ("2026-2027", 12, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(line)
        set_run_font(r, size=size, bold=bold)


def add_institution_pages(doc):
    doc.add_page_break()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("INSTITUTION AND DEPARTMENT VISION AND MISSION").bold = True
    institution = [
        ("VISION OF THE INSTITUTE", "To become one of the India’s leading engineering institutes in both education and research. We are committed to provide quality and state of the art technical education to our student, so that they become technologically superior and in turn contribute for creating a great society."),
        ("MISSION OF THE INSTITUTE", "Fostering a dynamic learning environment that equips students with Technical expertise, problem-solving skills and a deep commitment to ethical practices."),
        ("", "To cultivate a culture of innovation, incubation, research and entrepreneurship that drives technological advancements."),
        ("", "To uphold the spirit of mutual excellence while interacting with stake holders of our Institutional ecosystem."),
        ("", "Promoting lifelong learning, professional growth and ensuring holistic development of students and the well being of society."),
        ("VISION OF THE DEPARTMENT", "To develop center of excellence which will produce graduates with strong technical and industry-oriented knowledge needed to achieve distinction in the discipline."),
        ("MISSION OF THE DEPARTMENT", "To guide, motivate and enrich student’s knowledge and skills towards development of Industrial IoT."),
        ("", "To achieve the excellence in various domains like workforce, higher studies or lifelong learning."),
        ("", "To strengthen links between industry through partnership and collaborative development works."),
    ]
    for heading, text in institution:
        if heading:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(heading)
            set_run_font(r, size=11.5, bold=True)
        add_body_paragraph(doc, text, front=True)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PROGRAM OUTCOMES")
    set_run_font(r, size=16, bold=True)
    po = [
        ("PO1: Engineering Knowledge", "Apply knowledge of mathematics, natural science, computing, engineering fundamentals and an engineering specialization as specified in WK1 to WK4 respectively to develop to the solution of complex engineering problems."),
        ("PO2: Problem Analysis", "Identify, formulate, review research literature and analyze complex engineering problems reaching substantiated conclusions with consideration for sustainable development. (WK1 to WK4)"),
        ("PO3: Design/Development of Solutions", "Design creative solutions for complex engineering problems and design/develop systems/components/processes to meet identified needs with consideration for the public health and safety, whole-life cost, net zero carbon, culture, society and environment as required. (WK5)"),
        ("PO4: Conduct Investigations of Complex Problems", "Conduct investigations of complex engineering problems using research-based knowledge including design of experiments, modelling, analysis & interpretation of data to provide valid conclusions.(WK8)."),
        ("PO5: Engineering Tool Usage", "Create, select and apply appropriate techniques, resources and modern engineering & IT tools, including prediction and modelling recognizing their limitations to solve complex engineering problems. (WK2 and WK6)"),
        ("PO6: The Engineer and The World", "Analyze and evaluate societal and environmental aspects while solving complex engineering problems for its impact on sustainability with reference to economy, health, safety, legal framework, culture and environment. (WK1, WK5, and WK7)."),
        ("PO7: Ethics", "Apply ethical principles and commit to professional ethics, human values, diversity and inclusion; adhere to national & international laws. (WK9)"),
        ("PO8: Individual and Collaborative Team work", "Function effectively as an individual, and as a member or leader in diverse/multi-disciplinary teams."),
        ("PO9: Communication", "Communicate effectively and inclusively within the engineering community and society at large, such as being able to comprehend and write effective reports and design documentation, make effective presentations considering cultural, language, and learning differences"),
        ("PO10: Project Management and Finance", "Apply knowledge and understanding of engineering management principles and economic decision-making and apply these to one’s own work, as a member and leader in a team, and to manage projects and in multidisciplinary environments."),
        ("PO11: Life-Long Learning", "Recognize the need for, and have the preparation and ability for i) independent and life-long learning ii) adaptability to new and emerging technologies and iii) critical thinking in the broadest context of technological change. (WK8)"),
    ]
    for code, desc in po:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.03
        r = p.add_run(code + ": ")
        set_run_font(r, size=9.4, bold=True)
        r = p.add_run(desc)
        set_run_font(r, size=9.4)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PROGRAM EDUCATIONAL OBJECTIVES AND PROGRAM SPECIFIC OUTCOMES")
    set_run_font(r, size=14, bold=True)
    objectives = [
        "To provide graduates with a solid foundation in mathematical, scientific and engineering fundamentals required to formulate, analyze & solve complex Engineering problems.",
        "To create graduates with technical expertise to handle multi-disciplinary projects and with ability to excel in professional career and/or higher education.",
        "To Inculcate leadership qualities among graduates with ethical responsibilities and global security awareness.",
    ]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("PROGRAM EDUCATIONAL OBJECTIVES (PEOs)").bold = True
    for i, item in enumerate(objectives, 1):
        add_body_paragraph(doc, f"PEO {i}. {item}", front=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(15)
    p.add_run("PROGRAM SPECIFIC OUTCOMES (PSOs)").bold = True
    add_body_paragraph(doc, "PSO 1. The ability to develop problem solving skills through programming techniques using appropriate principles and concepts of Internet of Things.", front=True)
    add_body_paragraph(doc, "PSO 2. Develop the understanding of various IOT application development tools and its implementation for industry applications.", front=True)


def add_front_matter(doc, logo_bytes):
    add_cover(doc, logo_bytes)
    add_institution_pages(doc)

    doc.add_page_break()
    for text, size, bold in [
        ("PRIYADARSHINI COLLEGE OF ENGINEERING, NAGPUR", 15, True),
        ("(An Autonomous Institute Affiliated to Rashtrasant Tukdoji Maharaj Nagpur University, Nagpur)", 10, False),
        ("Department of Industrial IoT", 14, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(25)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CERTIFICATE")
    set_run_font(r, size=18, bold=True)
    for text in [
        "This is to certify that the project report entitled “Design and Implementation of a Digital Twin Orchestration Platform for Industrial IoT Manufacturing” is a bonafide work carried out by [Student Name(s)].",
        "The project report is submitted to Priyadarshini College of Engineering, Nagpur, an Autonomous Institute affiliated to Rashtrasant Tukdoji Maharaj Nagpur University, Nagpur, in partial fulfilment of the requirements for the degree of Bachelor of Technology in Industrial IoT for Session 2026-2027.",
    ]:
        add_body_paragraph(doc, text)
    for label in ["Project Guide", "Head of Department", "Principal", "Student Name(s)"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(22)
        p.add_run(f"{label}: ______________________________________________")

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DECLARATION")
    set_run_font(r, size=18, bold=True)
    add_body_paragraph(doc, "We, the undersigned, declare that the project entitled “Design and Implementation of a Digital Twin Orchestration Platform for Industrial IoT Manufacturing”, submitted in partial fulfilment of the requirements for the degree of Bachelor of Technology in Industrial IoT at Priyadarshini College of Engineering, Nagpur, is work carried out by us.")
    add_body_paragraph(doc, "We have identified the project sources and references used in this report. The implementation evidence and limitations are represented to the best of our knowledge. Any remaining institutional declaration requirements shall be completed and signed by the project team before submission.")
    for text in ["Place: Nagpur", "Date: ____________________", "Project team: ____________________", "Student signatures: ______________________________________________"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.add_run(text)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ACKNOWLEDGEMENT")
    set_run_font(r, size=18, bold=True)
    for text in [
        "We express our sincere gratitude to our project guide, [Project Guide Name], for guidance, technical feedback, and support throughout the design and implementation of this work.",
        "We thank the Department of Industrial IoT and the faculty of Priyadarshini College of Engineering for the academic environment and facilities that supported this project.",
        "We acknowledge the contributions of the project team members and all those who provided useful discussions, review, and encouragement. Their input helped refine the system architecture, implementation priorities, and presentation of the work.",
        "We also thank our families and friends for their patience and support during the project. Names and additional acknowledgements will be completed by the project team before submission.",
    ]:
        add_body_paragraph(doc, text)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ABSTRACT")
    set_run_font(r, size=18, bold=True)
    abstract = (
        "This project presents a simulator-backed Digital Twin Orchestration Platform for a discrete five-station automotive-component manufacturing line. "
        "The platform combines deterministic reduced-order models for Stamping, CNC machining, Welding, Inspection, and Packaging with a versioned event contract, a Live Twin state reducer, a process-level directed acyclic graph, production-token lineage, isolated scenario simulation, a FastAPI service, and a React dashboard. "
        "The design separates the model that generates observations from the service that materializes current twin state. Events carry stable identity, machine sequence, schema version, and timezone-aware timestamps. The reducer validates the payload, records observed receipt time, prevents duplicate writes from changing state, and preserves current state when a late sequence arrives. "
        "At the line level, the simulator links adjacent stations, enforces upstream starvation and a maximum downstream buffer of five units, and maintains the output conservation relation across the sequential route. Production tokens carry selected quality and process flags so that an inspection outcome can be traced to upstream model signals. "
        "The Scenario Worker captures a frozen line snapshot and runs baseline and counterfactual factories using the same seed. It computes comparative output, quality, energy, health, scrap, bottleneck, and trajectory fields and checks that the Live Twin state remains unchanged. The web application exposes machine, process, event, scenario, and system views. Docker Compose and Kubernetes resources describe a multi-service deployment. "
        "The project repository records a final backend verification summary of 48 tests with two skipped and a successful frontend production build. These results support software-behavior claims only. The station coefficients and cycle times are model configuration; physical calibration, real factory connectivity, measured latency benchmarks, and production improvement have not been established. The work provides a traceable prototype and an implementation basis for future hardware integration, security hardening, calibration, and empirical evaluation."
    )
    add_body_paragraph(doc, abstract)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("Keywords: ")
    set_run_font(r, size=11, bold=True)
    r = p.add_run("Industrial IoT, digital twin, manufacturing, MQTT, event processing, causal simulation, production flow, counterfactual analysis")
    set_run_font(r, size=11)

    doc.add_page_break()
    add_contents_pages(doc)


def configure_document(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.65)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.32)
        section.different_first_page_header_footer = True
        sectpr = section._sectPr
        old_border = sectpr.find(qn("w:pgBorders"))
        if old_border is not None:
            sectpr.remove(old_border)

    normal = doc.styles["Normal"]
    set_style_font(normal, size=12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    for name, size, align in [
        ("Title", 18, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 13, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 11.5, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 4", 10.5, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(style, size=size, bold=True)
        style.paragraph_format.alignment = align
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(11 if name != "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.0

    try:
        code_style = doc.styles["Code Listing"]
    except KeyError:
        code_style = doc.styles.add_style("Code Listing", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(code_style, name="Consolas", size=7.6)
    code_style.paragraph_format.line_spacing = Pt(9)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.left_indent = Inches(0.02)
    code_style.paragraph_format.keep_together = False

    try:
        caption_style = doc.styles["Caption"]
    except KeyError:
        caption_style = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(caption_style, size=9.5)
    caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_style.paragraph_format.keep_together = True
    caption_style.paragraph_format.space_before = Pt(3)
    caption_style.paragraph_format.space_after = Pt(7)

    section = doc.sections[0]
    footer = section.footer
    add_page_number(footer.paragraphs[0])
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""

    # Ensure Word refreshes table and page fields when opened.
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_markdown_content(doc, logo_bytes):
    add_front_matter(doc, logo_bytes)
    figure_counts = defaultdict(int)
    current_chapter = 1
    lines = CONTENT.read_text(encoding="utf-8-sig").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                if not all(set(v.replace(":", "")) <= {"-"} for v in values):
                    table_rows.append(values)
                index += 1
            add_md_table(doc, table_rows)
            continue
        if line.startswith("[[FIGURE:"):
            name = line[len("[[FIGURE:"):-2]
            add_figure(doc, name, current_chapter, figure_counts)
            index += 1
            continue
        if line == "[[SCREENSHOT]]":
            add_figure(doc, "screenshot", current_chapter, figure_counts)
            index += 1
            continue
        if line.startswith("[[SOURCE:"):
            name = line[len("[[SOURCE:"):-2]
            add_source_listing(doc, ROOT / name)
            index += 1
            continue
        if line.startswith("# "):
            heading = line[2:].strip()
            chapter_match = re.match(r"Chapter\s+(\d+)", heading, flags=re.I)
            if chapter_match:
                current_chapter = int(chapter_match.group(1))
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.page_break_before = True
            p.add_run(heading.upper())
            index += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(line[3:].strip())
            index += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            p.add_run(line[4:].strip())
            index += 1
            continue
        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.space_after = Pt(3)
            p.add_run("• " + line[2:].replace("**", ""))
            index += 1
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.add_run(line.replace("**", ""))
            index += 1
            continue
        # Gather a complete prose paragraph until the next structural token.
        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            upcoming = lines[index].rstrip()
            if (not upcoming or upcoming.startswith("#") or upcoming.startswith("|")
                    or upcoming.startswith("[[") or upcoming.startswith("- ")
                    or upcoming.startswith("* ") or re.match(r"^\d+\.\s", upcoming)):
                break
            paragraph_lines.append(upcoming)
            index += 1
        add_body_paragraph(doc, " ".join(paragraph_lines))


def add_source_listing(doc, path: Path):
    if not path.exists():
        raise FileNotFoundError(path)
    doc.add_page_break()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    r = title.add_run(f"Source listing: {path.relative_to(ROOT).as_posix()}")
    set_run_font(r, size=9.5, bold=True)
    source_lines = path.read_text(encoding="utf-8-sig").splitlines()
    physical_lines: list[str] = []
    for line_no, line in enumerate(source_lines, 1):
        prefix = f"{line_no:04d} "
        if not line:
            physical_lines.append(prefix)
            continue
        wrapped = textwrap.wrap(
            line, width=103, subsequent_indent="     ",
            break_long_words=True, break_on_hyphens=False,
            replace_whitespace=False, drop_whitespace=False,
        )
        physical_lines.extend(prefix + wrapped[0] if i == 0 else part for i, part in enumerate(wrapped))
    chunk_size = 70 if path.name == "mongo_scenario_repository.py" else 64
    for chunk_index, offset in enumerate(range(0, len(physical_lines), chunk_size)):
        if chunk_index:
            doc.add_page_break()
            label = doc.add_paragraph()
            label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            label.paragraph_format.space_after = Pt(2)
            r = label.add_run(f"{path.name} - lines {offset+1}-{min(offset+chunk_size, len(physical_lines))}")
            set_run_font(r, size=8.5, italic=True, color="555555")
        p = doc.add_paragraph(style="Code Listing")
        p.paragraph_format.keep_together = False
        run = p.add_run("\n".join(physical_lines[offset:offset + chunk_size]))
        set_run_font(run, name="Consolas", size=7.6)


def remove_sample_photo_relationship(doc):
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/image") and rel.target_ref.endswith("image2.jpeg"):
            del doc.part.rels[rid]


def main():
    verify_template()
    create_diagrams()
    with zipfile.ZipFile(TEMPLATE) as archive:
        logo_bytes = archive.read("word/media/image1.jpeg")

    doc = Document(str(TEMPLATE))
    remove_sample_photo_relationship(doc)
    body = doc._element.body
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    configure_document(doc)
    add_markdown_content(doc, logo_bytes)
    doc.core_properties.title = "Design and Implementation of a Digital Twin Orchestration Platform for Industrial IoT Manufacturing"
    doc.core_properties.subject = "Bachelor of Technology project thesis, Industrial IoT, 2026-2027"
    doc.core_properties.keywords = "Industrial IoT, digital twin, manufacturing, thesis"
    doc.core_properties.author = "[Student Name(s)]"
    doc.save(str(OUTPUT))
    print(f"Saved {OUTPUT}")
    print(f"Content source: {CONTENT}")
    print(f"Generated figures: {len(list(FIGURES.glob('*.png')))}")


if __name__ == "__main__":
    main()
